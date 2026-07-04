#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export generated equivalent-model building parameters to Excel."""
from __future__ import annotations

import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(PROJECT_DIR))

from models.build_arranger import arrange_buildings
from models.combustibles import SPECIALIZED_COMPONENTS
from models.facility import FacilityManager
from models.geometry import detect_coplanar_openings, is_coplanar
from models.materials import COMBUSTIBLE_LIBRARY


SCALES = (
    ("small", 0, "小"),
    ("medium", 1, "中"),
    ("large", 2, "大"),
)

WALLS = ("x_min", "x_max", "y_min", "y_max")

EXPORT_GRID_BY_FACILITY = {
    "metallurgical_facilities": {"small": 1.0, "medium": 2.0, "large": 2.0},
    "machinery_manufacturing": {"small": 1.0, "medium": 2.0, "large": 3.0},
    "airport_hangar": {"small": 1.0, "medium": 1.0, "large": 1.0},
    "aerospace": {"small": 3.0, "medium": 3.0, "large": 3.0},
}


def _opening_area(opening) -> float:
    return float(opening.boundary[1]) * float(opening.boundary[3])


def _counter_summary(counter: Counter) -> str:
    if not counter:
        return ""
    return "; ".join(f"{key}:{value}" for key, value in sorted(counter.items()))


def _count_entries(entries: Iterable[dict]) -> tuple[int, Counter]:
    total = 0
    by_key: Counter = Counter()
    for entry in entries:
        count = int(entry.get("count", 0) or 0)
        key = entry.get("key", "")
        total += count
        if key:
            by_key[key] += count
    return total, by_key


def _content_rows(facility, facility_cn, scale, scale_cn, building, story, scope, zone, entries, kind):
    rows = []
    for entry in entries:
        key = entry.get("key", "")
        count = int(entry.get("count", 0) or 0)
        if not key or count <= 0:
            continue

        if kind == "combustible":
            definition = COMBUSTIBLE_LIBRARY.get(key, {})
            name = definition.get("name", key)
            category = "combustible"
            length = definition.get("length")
            width = definition.get("width")
            height = definition.get("height")
            item_volume = (
                float(length) * float(width) * float(height)
                if length is not None and width is not None and height is not None
                else None
            )
        else:
            comp = SPECIALIZED_COMPONENTS.get(key)
            name = comp.name if comp else key
            category = comp.category if comp else "component"
            length = comp.total_length if comp else None
            width = comp.total_width if comp else None
            height = comp.total_height if comp else None
            item_volume = (
                float(length) * float(width) * float(height)
                if length is not None and width is not None and height is not None
                else None
            )

        rows.append({
            "设施": facility,
            "设施中文名": facility_cn,
            "规模": scale,
            "规模中文": scale_cn,
            "建筑": building.name,
            "建筑中文名": building.cn_name,
            "楼层": story.name,
            "范围": scope,
            "区域": zone,
            "类型": kind,
            "key": key,
            "名称": name,
            "类别": category,
            "数量": count,
            "单体长_m": length,
            "单体宽_m": width,
            "单体高_m": height,
            "单体包络体积_m3": item_volume,
            "总包络体积_m3": item_volume * count if item_volume is not None else None,
            "size_multiplier": entry.get("size_multiplier", ""),
            "orientation": entry.get("orientation", ""),
        })
    return rows


def _facility_extent(buildings) -> dict:
    if not buildings:
        return {
            "facility_x_min": 0,
            "facility_x_max": 0,
            "facility_y_min": 0,
            "facility_y_max": 0,
            "facility_footprint_m2": 0,
        }
    x0 = min(b.offset_x for b in buildings)
    x1 = max(b.offset_x + b.length for b in buildings)
    y0 = min(b.offset_y for b in buildings)
    y1 = max(b.offset_y + b.width for b in buildings)
    return {
        "facility_x_min": x0,
        "facility_x_max": x1,
        "facility_y_min": y0,
        "facility_y_max": y1,
        "facility_footprint_m2": (x1 - x0) * (y1 - y0),
    }


def _load_generated_buildings(mgr: FacilityManager, facility: str, scale_idx: int):
    buildings = []
    params_by_building = {}
    for building_name in mgr.list_buildings(facility):
        params = mgr.params_for_scale(facility, building_name, scale_idx=scale_idx)
        params_by_building[building_name] = params
        buildings.append(mgr.load_equivalent(facility, building_name, params))
    if len(buildings) > 1:
        layout_applied = mgr.arrange_buildings_by_layout(facility, buildings)
        if not layout_applied:
            arrange_buildings(buildings, gap=20.0)
    for building in buildings:
        building.update_z_offsets()
    return buildings, params_by_building


def collect_rows():
    mgr = FacilityManager()
    facilities = [name for name in sorted(mgr.facilities) if mgr.get_type(name) == "equivalent"]

    summary_rows = []
    wall_rows = []
    story_rows = []
    fc_rows = []
    content_rows = []
    facility_scale_rows = []

    for facility in facilities:
        facility_data = mgr.facilities[facility]
        facility_cn = facility_data.get("cn_name", facility_data.get("name", ""))
        for scale, scale_idx, scale_cn in SCALES:
            buildings, params_by_building = _load_generated_buildings(mgr, facility, scale_idx)
            extent = _facility_extent(buildings)
            export_grid = EXPORT_GRID_BY_FACILITY.get(facility, {}).get(scale, "")

            facility_scale_rows.append({
                "设施": facility,
                "设施中文名": facility_cn,
                "规模": scale,
                "规模中文": scale_cn,
                "建筑数量": len(buildings),
                "导出网格_m": export_grid,
                "设施包络_x_min": extent["facility_x_min"],
                "设施包络_x_max": extent["facility_x_max"],
                "设施包络_y_min": extent["facility_y_min"],
                "设施包络_y_max": extent["facility_y_max"],
                "设施包络面积_m2": extent["facility_footprint_m2"],
                "建筑名列表": "; ".join(b.name for b in buildings),
            })

            for building in buildings:
                params = params_by_building[building.name]
                floor_area = building.length * building.width * max(1, len(building.stories))
                volume = building.length * building.width * building.height
                story_heights = [story.height for story in building.stories]
                wall_stats = {
                    wall: {
                        "door_count": 0,
                        "window_count": 0,
                        "other_count": 0,
                        "total_count": 0,
                        "door_area": 0.0,
                        "window_area": 0.0,
                        "total_area": 0.0,
                    }
                    for wall in WALLS
                }
                exterior_openings_total = 0
                interior_openings_total = 0
                exterior_door_area = 0.0
                exterior_window_area = 0.0
                roof_opening_count = 0
                fc_count = 0
                fc_area_total = 0.0
                combustible_total = 0
                component_total = 0
                combustible_by_key: Counter = Counter()
                component_by_key: Counter = Counter()

                for story_index, story in enumerate(building.stories, 1):
                    exterior_openings = list(story.openings) + detect_coplanar_openings(building, story)
                    story_wall_counter = Counter()
                    story_door_count = 0
                    story_window_count = 0
                    story_other_count = 0
                    story_opening_area = 0.0
                    for opening in exterior_openings:
                        wall = opening.wall
                        if wall not in wall_stats:
                            continue
                        area = _opening_area(opening)
                        typ = opening.type
                        exterior_openings_total += 1
                        story_opening_area += area
                        story_wall_counter[(wall, typ)] += 1
                        wall_stats[wall]["total_count"] += 1
                        wall_stats[wall]["total_area"] += area
                        if typ == "door":
                            wall_stats[wall]["door_count"] += 1
                            wall_stats[wall]["door_area"] += area
                            story_door_count += 1
                            exterior_door_area += area
                        elif typ == "window":
                            wall_stats[wall]["window_count"] += 1
                            wall_stats[wall]["window_area"] += area
                            story_window_count += 1
                            exterior_window_area += area
                        else:
                            wall_stats[wall]["other_count"] += 1
                            story_other_count += 1

                    story_cb_total, story_cb_by_key = _count_entries(story.combustibles)
                    story_sc_total, story_sc_by_key = _count_entries(story.specialized_components)
                    combustible_total += story_cb_total
                    component_total += story_sc_total
                    combustible_by_key.update(story_cb_by_key)
                    component_by_key.update(story_sc_by_key)
                    roof_opening_count += len(story.roof.openings)

                    content_rows.extend(_content_rows(
                        facility, facility_cn, scale, scale_cn, building, story,
                        "story", story.name, story.combustibles, "combustible",
                    ))
                    content_rows.extend(_content_rows(
                        facility, facility_cn, scale, scale_cn, building, story,
                        "story", story.name, story.specialized_components, "component",
                    ))

                    story_rows.append({
                        "设施": facility,
                        "设施中文名": facility_cn,
                        "规模": scale,
                        "规模中文": scale_cn,
                        "建筑": building.name,
                        "建筑中文名": building.cn_name,
                        "楼层序号": story_index,
                        "楼层": story.name,
                        "层高_m": story.height,
                        "z_bottom_m": story.z_bottom,
                        "z_top_m": story.z_top,
                        "建筑长_m": building.length,
                        "建筑宽_m": building.width,
                        "楼层面积_m2": building.length * building.width,
                        "外墙门数量": story_door_count,
                        "外墙窗数量": story_window_count,
                        "外墙其它洞口数量": story_other_count,
                        "外墙洞口总数": len(exterior_openings),
                        "外墙洞口面积_m2": story_opening_area,
                        "防火分区数量": len(story.fire_compartments),
                        "屋顶洞口数量": len(story.roof.openings),
                        "楼层可燃物数量": story_cb_total,
                        "楼层组件数量": story_sc_total,
                        "楼层可燃物key汇总": _counter_summary(story_cb_by_key),
                        "楼层组件key汇总": _counter_summary(story_sc_by_key),
                        "方向洞口汇总": _counter_summary(story_wall_counter),
                    })

                    for fc_index, fc in enumerate(story.fire_compartments, 1):
                        fc_count += 1
                        x_min, x_max, y_min, y_max = fc.boundary
                        fc_length = x_max - x_min
                        fc_width = y_max - y_min
                        fc_area = fc_length * fc_width
                        fc_area_total += fc_area

                        fc_ext_counter = Counter()
                        fc_int_counter = Counter()
                        fc_ext_area = 0.0
                        fc_int_area = 0.0
                        for opening in fc.openings:
                            area = _opening_area(opening)
                            if is_coplanar(fc.boundary, building.length, building.width, opening.wall):
                                fc_ext_counter[(opening.wall, opening.type)] += 1
                                fc_ext_area += area
                            else:
                                fc_int_counter[(opening.wall, opening.type)] += 1
                                fc_int_area += area
                                interior_openings_total += 1

                        fc_cb_total, fc_cb_by_key = _count_entries(fc.combustibles)
                        fc_sc_total, fc_sc_by_key = _count_entries(fc.specialized_components)
                        combustible_total += fc_cb_total
                        component_total += fc_sc_total
                        combustible_by_key.update(fc_cb_by_key)
                        component_by_key.update(fc_sc_by_key)

                        content_rows.extend(_content_rows(
                            facility, facility_cn, scale, scale_cn, building, story,
                            "fire_compartment", fc.name, fc.combustibles, "combustible",
                        ))
                        content_rows.extend(_content_rows(
                            facility, facility_cn, scale, scale_cn, building, story,
                            "fire_compartment", fc.name, fc.specialized_components, "component",
                        ))

                        fc_rows.append({
                            "设施": facility,
                            "设施中文名": facility_cn,
                            "规模": scale,
                            "规模中文": scale_cn,
                            "建筑": building.name,
                            "建筑中文名": building.cn_name,
                            "楼层": story.name,
                            "分区序号": fc_index,
                            "防火分区": fc.name,
                            "x_min_m": x_min,
                            "x_max_m": x_max,
                            "y_min_m": y_min,
                            "y_max_m": y_max,
                            "分区长_m": fc_length,
                            "分区宽_m": fc_width,
                            "分区面积_m2": fc_area,
                            "占建筑单层面积比例": fc_area / (building.length * building.width) if building.length and building.width else 0,
                            "防火墙厚度_m": fc.firewall_thickness,
                            "防火墙材料": fc.firewall_material,
                            "外墙共面门窗数量": sum(fc_ext_counter.values()),
                            "内部墙门窗数量": sum(fc_int_counter.values()),
                            "外墙共面门窗面积_m2": fc_ext_area,
                            "内部墙门窗面积_m2": fc_int_area,
                            "外墙共面门窗汇总": _counter_summary(fc_ext_counter),
                            "内部墙门窗汇总": _counter_summary(fc_int_counter),
                            "可燃物数量": fc_cb_total,
                            "组件数量": fc_sc_total,
                            "可燃物key汇总": _counter_summary(fc_cb_by_key),
                            "组件key汇总": _counter_summary(fc_sc_by_key),
                        })

                for wall, stats in wall_stats.items():
                    wall_rows.append({
                        "设施": facility,
                        "设施中文名": facility_cn,
                        "规模": scale,
                        "规模中文": scale_cn,
                        "建筑": building.name,
                        "建筑中文名": building.cn_name,
                        "方向": wall,
                        "门数量": stats["door_count"],
                        "窗数量": stats["window_count"],
                        "其它洞口数量": stats["other_count"],
                        "门窗洞口总数": stats["total_count"],
                        "门面积_m2": stats["door_area"],
                        "窗面积_m2": stats["window_area"],
                        "洞口总面积_m2": stats["total_area"],
                    })

                summary = {
                    "设施": facility,
                    "设施中文名": facility_cn,
                    "规模": scale,
                    "规模中文": scale_cn,
                    "建筑": building.name,
                    "建筑中文名": building.cn_name,
                    "导出网格_m": export_grid,
                    "长_m": building.length,
                    "宽_m": building.width,
                    "高_m": building.height,
                    "层数": len(building.stories),
                    "平均层高_m": sum(story_heights) / len(story_heights) if story_heights else 0,
                    "层高列表_m": "; ".join(f"{h:g}" for h in story_heights),
                    "占地面积_m2": building.length * building.width,
                    "总楼面面积_m2": floor_area,
                    "包络体积_m3": volume,
                    "墙厚_m": building.wall_thickness,
                    "offset_x_m": building.offset_x,
                    "offset_y_m": building.offset_y,
                    "参数length_m": params.get("length"),
                    "参数width_m": params.get("width"),
                    "参数height_m": params.get("height"),
                    "参数stories": params.get("stories"),
                    "防火分区数量": fc_count,
                    "防火分区总面积_m2": fc_area_total,
                    "外墙门数量": sum(wall_stats[w]["door_count"] for w in WALLS),
                    "外墙窗数量": sum(wall_stats[w]["window_count"] for w in WALLS),
                    "外墙其它洞口数量": sum(wall_stats[w]["other_count"] for w in WALLS),
                    "外墙门窗洞口总数": exterior_openings_total,
                    "内部墙门窗数量": interior_openings_total,
                    "外墙门面积_m2": exterior_door_area,
                    "外墙窗面积_m2": exterior_window_area,
                    "外墙洞口总面积_m2": sum(wall_stats[w]["total_area"] for w in WALLS),
                    "屋顶洞口数量": roof_opening_count,
                    "可燃物总数量": combustible_total,
                    "组件总数量": component_total,
                    "可燃物key汇总": _counter_summary(combustible_by_key),
                    "组件key汇总": _counter_summary(component_by_key),
                }
                for wall in WALLS:
                    summary[f"{wall}_门数量"] = wall_stats[wall]["door_count"]
                    summary[f"{wall}_窗数量"] = wall_stats[wall]["window_count"]
                    summary[f"{wall}_其它洞口数量"] = wall_stats[wall]["other_count"]
                    summary[f"{wall}_门窗洞口总数"] = wall_stats[wall]["total_count"]
                summary_rows.append(summary)

    return {
        "建筑汇总": summary_rows,
        "方向门窗": wall_rows,
        "楼层": story_rows,
        "防火分区": fc_rows,
        "内容物": content_rows,
        "设施规模": facility_scale_rows,
    }


def _write_sheet(wb: Workbook, name: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title=name)
    if not rows:
        ws.append(["无数据"])
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        header = str(column_cells[0].value or "")
        max_len = len(header)
        for cell in column_cells[1:]:
            value = cell.value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 42)


def export(path: Path) -> Path:
    rows_by_sheet = collect_rows()
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in ("建筑汇总", "方向门窗", "楼层", "防火分区", "内容物", "设施规模"):
        _write_sheet(wb, sheet_name, rows_by_sheet[sheet_name])
    wb.save(path)
    return path


def main() -> int:
    out_path = PROJECT_DIR / "export" / "equivalent_model_building_parameters.xlsx"
    if len(sys.argv) > 1:
        out_path = Path(sys.argv[1]).expanduser()
        if not out_path.is_absolute():
            out_path = PROJECT_DIR / out_path
    export(out_path)
    print(out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
